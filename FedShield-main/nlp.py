import re
import nltk
from nltk.corpus import stopwords
import os
import pandas as pd
import openpyxl as xl
from docx import Document

stop = stopwords.words('english')

def extract_phone_numbers(string):
    r = re.compile(r'(\d{3}[-\.\s]??\d{3}[-\.\s]??\d{5}|\(\d{3}\)\s*\d{3}[-\.\s]??\d{5}|\d{3}[-\.\s]??\d{5})')
    phone_numbers = r.findall(string)
    return [re.sub(r'\D', '', number) for number in phone_numbers]

def extract_CNIC(string):
    r = re.compile(r'\d{13}|\d{5}-\d{7}-\d{1}')
    CNIC=r.findall(string)
    return [re.sub(r'\D', '', number) for number in CNIC]

def extract_PAnum(string):
    r = re.compile(r'PA\d{5}|PA\s*\d{5}')
    PAnum=r.findall(string)
    return PAnum

def extract_email_addresses(string):
    r = re.compile(r'[\w\.-]+@[\w\.-]+.com')
    return r.findall(string)


def ie_preprocess(document):
    document = ' '.join([i for i in document.split() if i not in stop])
    sentences = nltk.sent_tokenize(document)
    sentences = [nltk.word_tokenize(sent) for sent in sentences]
    sentences = [nltk.pos_tag(sent) for sent in sentences]
    return sentences


def extract_names(document):
    names = []
    sentences = ie_preprocess(document)
    for tagged_sentence in sentences:
        for chunk in nltk.ne_chunk(tagged_sentence):
            if type(chunk) == nltk.tree.Tree:
                if chunk.label() == 'PERSON':
                    names.append(' '.join([c[0] for c in chunk]))
    return names

def nlp_check(filepath):
    name = os.path.basename(os.path.basename(filepath))
    _, extension = os.path.splitext(name)
    if (extension == ".txt" ):
        with open(filepath, 'r') as file:
            data = file.read()
            print(data)
        numbers = extract_phone_numbers(data)
        PAnum = extract_PAnum(data)
        CNIC = extract_CNIC(data)
        if numbers or PAnum or CNIC:
            return 0
        else:
            return 1
    elif extension==".docx":
        doc = Document(filepath)
        for para in doc.paragraphs:
            string=para.text
            numbers = extract_phone_numbers(string)
            PAnum= extract_PAnum(string)
            CNIC=extract_CNIC(string)
            if numbers or PAnum or CNIC:
                return 0
        return 1
    elif extension==".xlsx":
        compare_wb = xl.load_workbook(filepath)
        compare_sheet_names = compare_wb.sheetnames

        for compare_sheet_name in compare_sheet_names:
            compare_df = pd.read_excel(filepath, sheet_name=compare_sheet_name)
            for column in compare_df.columns:
                compare_keywords = compare_df[column].dropna().tolist()
                my_string = ' '.join(compare_keywords)
                numbers = extract_phone_numbers(my_string)
                PAnum = extract_PAnum(my_string)
                CNIC = extract_CNIC(my_string)
                if numbers or PAnum or CNIC:
                    return 0
        return 1
